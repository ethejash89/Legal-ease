from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF


def _safe_filename(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9_-]+", "_", value.strip())
    return value.strip("_") or "legal_document"


def export_txt(content: str) -> bytes:
    return content.encode("utf-8")


def export_docx(
    content: str,
    document_type: str,
    parties: str = "",
    effective_date: str = "",
) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(document_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    if effective_date:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"Effective Date: {effective_date}")
        r.italic = True

    if parties:
        p = doc.add_paragraph()
        r = p.add_run("Parties: ")
        r.bold = True
        p.add_run(parties)

    doc.add_paragraph()

    for raw in content.splitlines():
        line = raw.strip()
        if not line:
            doc.add_paragraph()
            continue

        p = doc.add_paragraph()
        if re.match(r"^\d+[\.\)]\s+", line):
            r = p.add_run(line)
            r.bold = True
        elif len(line) < 100 and line.upper() == line and any(c.isalpha() for c in line):
            r = p.add_run(line)
            r.bold = True
            r.font.size = Pt(12)
        else:
            p.add_run(line)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase - AI-assisted draft | Review before execution").font.size = Pt(8)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


class BrandedPDF(FPDF):
    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", size=8)
        self.cell(0, 8, "LegalEase - AI-assisted draft | Review before execution", align="C")


def export_pdf(content: str, document_type: str) -> bytes:
    pdf = BrandedPDF()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, document_type.upper(), ln=True, align="C")
    pdf.ln(5)

    for raw in content.splitlines():
        line = raw.strip()
        if not line:
            pdf.ln(4)
            continue

        if re.match(r"^\d+[\.\)]\s+", line) or (
            len(line) < 100 and line.upper() == line and any(c.isalpha() for c in line)
        ):
            pdf.set_font("Helvetica", "B", 11)
        else:
            pdf.set_font("Helvetica", size=10.5)

        # fpdf2's multi_cell handles wrapping for long legal clauses.
        pdf.multi_cell(pdf.epw, 6, line)

    return bytes(pdf.output())
